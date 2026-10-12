"""Build real documents in memory for tests: PDF, Word, Excel, PowerPoint, zips.

Using real files made by the real libraries (not hand-written bytes) means the parsers are tested
on what Word, Excel and PowerPoint style documents actually look like inside.
"""

import io
import zipfile
from collections.abc import Callable, Sequence

import openpyxl
from docx import Document as new_word_document
from docx.document import Document as WordDocument
from pptx import Presentation as new_presentation
from pptx.presentation import Presentation
from pptx.util import Inches
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

STYLES = getSampleStyleSheet()


def zip_bytes(
    entries: dict[str, bytes], compression: int = zipfile.ZIP_DEFLATED
) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    return buffer.getvalue()


def add_to_zip(original: bytes, extra: dict[str, bytes]) -> bytes:
    """A copy of a zip file with extra entries added."""
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(original)) as src,
        zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst,
    ):
        for info in src.infolist():
            dst.writestr(info, src.read(info))
        for name, data in extra.items():
            dst.writestr(name, data)
    return out.getvalue()


# --- PDF --------------------------------------------------------------------------------


def pdf_text(*lines: str) -> bytes:
    """A one-page PDF with a few lines of text."""
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    y = 800
    for line in lines or ("Hello world",):
        pdf.drawString(72, y, line)
        y -= 18
    pdf.save()
    return buffer.getvalue()


def pdf_flowables(flowables: list, margins: float = 18 * mm) -> bytes:
    buffer = io.BytesIO()
    SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=margins,
        rightMargin=margins,
        topMargin=margins,
        bottomMargin=margins,
    ).build(flowables)
    return buffer.getvalue()


def para(text: str, style: str = "BodyText") -> Paragraph:
    return Paragraph(text, STYLES[style])


def grid_table(
    rows: list[list[str]],
    repeat_header: bool = False,
    widths: Sequence[float] | None = None,
) -> Table:
    table = Table(rows, colWidths=widths, repeatRows=1 if repeat_header else 0)
    table.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                ("FONT", (0, 0), (-1, -1), "Helvetica", 9),
            ]
        )
    )
    return table


def borderless_table(
    rows: list[list[str]],
    widths: Sequence[float] | None = None,
    total_rows: Sequence[int] = (),
) -> Table:
    """A financial-statement style table: no borders, numbers right-aligned, rules over totals."""
    table = Table(rows, colWidths=widths)
    style = [
        ("FONT", (0, 0), (-1, -1), "Helvetica", 9),
        ("ALIGN", (1, 0), (-1, -1), "RIGHT"),
        ("LINEBELOW", (0, 0), (-1, 0), 0.6, colors.black),
    ]
    style += [("LINEABOVE", (0, r), (-1, r), 0.4, colors.black) for r in total_rows]
    table.setStyle(TableStyle(style))
    return table


def pdf_scanned() -> bytes:
    """An image-only PDF page (no text layer), like a scan."""
    from PIL import Image

    image = Image.effect_noise((600, 800), 40).convert("RGB")
    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    png = io.BytesIO()
    image.save(png, format="PNG")
    png.seek(0)
    from reportlab.lib.utils import ImageReader

    pdf.drawImage(ImageReader(png), 0, 0, width=A4[0], height=A4[1])
    pdf.save()
    return buffer.getvalue()


def pdf_page_break() -> PageBreak:
    return PageBreak()


# --- Word -------------------------------------------------------------------------------


def docx_bytes(build: Callable[[WordDocument], None] | None = None) -> bytes:
    document = new_word_document()
    if build is None:
        document.add_paragraph("Hello from Word.")
    else:
        build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --- Excel ------------------------------------------------------------------------------


def xlsx_bytes(build: Callable[[openpyxl.Workbook], None] | None = None) -> bytes:
    workbook = openpyxl.Workbook()
    if build is None:
        workbook.active.append(["name", "value"])  # type: ignore[union-attr]
        workbook.active.append(["a", 1])  # type: ignore[union-attr]
    else:
        build(workbook)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --- PowerPoint -------------------------------------------------------------------------


def pptx_bytes(build: Callable[[Presentation], None] | None = None) -> bytes:
    presentation = new_presentation()
    if build is None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Title"  # type: ignore[union-attr]
        slide.placeholders[1].text = "Body"
    else:
        build(presentation)
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def inches(value: float) -> int:
    return Inches(value)


def spacer(height: float = 6) -> Spacer:
    return Spacer(1, height)
