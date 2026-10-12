"""PowerPoint (.pptx) parser.

A slide is parsed in reading order (top to bottom, left to right). Titles become headings,
tables keep their merged cells, speaker notes are kept and marked, and chart data that lives
inside the presentation (a real embedded workbook, not pixels) is read as a table.
"""

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from backend.app.config import Settings
from backend.app.ingest.builder import BlockBuilder
from backend.app.ingest.models import PageReport, ParsedDocument, ParseError

_BODY_PLACEHOLDERS = {
    PP_PLACEHOLDER.BODY,
    PP_PLACEHOLDER.OBJECT,
    PP_PLACEHOLDER.SUBTITLE,
}
_MAX_SHAPE_DEPTH = 20


def _is_hidden(shape: object) -> bool:
    element = getattr(shape, "_element", None)
    if element is None:
        return False
    hidden = element.xpath(".//p:cNvPr/@hidden")
    return bool(hidden) and str(hidden[0]) in {"1", "true"}


def _ordered(shapes: object) -> list[object]:
    return sorted(
        shapes, key=lambda s: (getattr(s, "top", 0) or 0, getattr(s, "left", 0) or 0)
    )  # type: ignore[call-overload]


class _SlideReader:
    def __init__(self, builder: BlockBuilder, slide_number: int) -> None:
        self.builder = builder
        self.slide = slide_number
        self.hidden_shapes = 0
        self.images = 0
        self.counter = 0

    def index(self) -> int:
        self.counter += 1
        return self.counter

    def shapes(self, shapes: object, depth: int = 0) -> None:
        if depth > _MAX_SHAPE_DEPTH:
            raise ParseError(
                "too_deeply_grouped",
                "The presentation nests shapes too deeply to read.",
            )
        for shape in _ordered(shapes):
            if _is_hidden(shape):
                self.hidden_shapes += 1
                continue
            kind = getattr(shape, "shape_type", None)
            if kind == MSO_SHAPE_TYPE.GROUP:
                self.shapes(shape.shapes, depth + 1)  # type: ignore[attr-defined]
            elif getattr(shape, "has_table", False) and shape.has_table:  # type: ignore[attr-defined]
                self.table(shape.table)  # type: ignore[attr-defined]
            elif getattr(shape, "has_chart", False) and shape.has_chart:  # type: ignore[attr-defined]
                self.chart(shape.chart)  # type: ignore[attr-defined]
            elif kind == MSO_SHAPE_TYPE.PICTURE:
                self.images += 1
            elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:  # type: ignore[attr-defined]
                self.text_frame(shape)

    def text_frame(self, shape: object) -> None:
        placeholder = (
            getattr(shape, "placeholder_format", None)
            if getattr(shape, "is_placeholder", False)
            else None
        )
        if placeholder is not None and placeholder.type in {
            PP_PLACEHOLDER.TITLE,
            PP_PLACEHOLDER.CENTER_TITLE,
        }:
            return  # the slide title was already used as the heading
        is_body = placeholder is not None and placeholder.type in _BODY_PLACEHOLDERS
        for paragraph in shape.text_frame.paragraphs:  # type: ignore[attr-defined]
            text = "".join(run.text for run in paragraph.runs) or paragraph.text
            if is_body or paragraph.level > 0:
                self.builder.list_item(
                    text, level=paragraph.level, slide=self.slide, index=self.index()
                )
            else:
                self.builder.paragraph(text, slide=self.slide, index=self.index())

    def table(self, table: object) -> None:
        rows: list[list[str]] = []
        merged = False
        for row in table.rows:  # type: ignore[attr-defined]
            out: list[str] = []
            for cell in row.cells:
                if cell.is_spanned:
                    out.append("")
                    merged = True
                else:
                    merged = merged or cell.is_merge_origin
                    out.append(cell.text)
            rows.append(out)
        previous = self.builder.blocks[-1].text[:300] if self.builder.blocks else ""
        self.builder.table(
            rows,
            has_merged_cells=merged,
            units_from=f"{previous} {' '.join(rows[0]) if rows else ''}",
            slide=self.slide,
            index=self.index(),
        )

    def chart(self, chart: object) -> None:
        """Chart data stored inside the file is real numbers, so it is read as a table."""
        try:
            title = chart.chart_title.text_frame.text if chart.has_title else ""  # type: ignore[attr-defined]
            plot = chart.plots[0]  # type: ignore[attr-defined]
            categories = [str(c) for c in plot.categories]
            series = list(plot.series)
        except IndexError, KeyError, AttributeError, ValueError:
            self.builder.figure(
                "",
                flags=["chart_data_unreadable"],
                slide=self.slide,
                index=self.index(),
            )
            return
        if title:
            self.builder.paragraph(title, slide=self.slide, index=self.index())
        rows = [["", *[str(s.name) for s in series]]]
        for i, category in enumerate(categories):
            rows.append([category, *[_number(list(s.values), i) for s in series]])
        self.builder.table(
            rows,
            flags=["chart_data_from_presentation"],
            slide=self.slide,
            index=self.index(),
        )


def _number(values: list[object], index: int) -> str:
    if index >= len(values) or values[index] is None:
        return ""
    value = values[index]
    return f"{value:g}" if isinstance(value, float) else str(value)


def parse_pptx(path: Path, settings: Settings) -> ParsedDocument:
    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        raise ParseError(
            "corrupt_pptx", "The PowerPoint file could not be opened."
        ) from exc

    builder = BlockBuilder()
    pages: list[PageReport] = []
    flags: list[str] = []
    for number, slide in enumerate(presentation.slides, 1):
        title_shape = slide.shapes.title
        title = (
            title_shape.text_frame.text.strip()
            if title_shape is not None and title_shape.has_text_frame
            else ""
        )
        before = len(builder.blocks)
        if title:
            builder.heading(title, 1, slide=number, index=0)
        reader = _SlideReader(builder, number)
        reader.shapes(slide.shapes)

        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text
            block = builder.paragraph(notes, slide=number, index=reader.index())
            if block is not None:
                block.flags.append("speaker_notes")

        page_flags: list[str] = []
        if slide._element.get("show") == "0":
            page_flags.append("hidden_slide")
        if reader.hidden_shapes:
            page_flags.append(f"hidden_shapes_removed:{reader.hidden_shapes}")
        if reader.images:
            page_flags.append(f"images_not_read:{reader.images}")
        added = builder.blocks[before:]
        pages.append(
            PageReport(
                index=number,
                label=title or None,
                status="ok" if added else "empty",
                chars=sum(len(b.text) for b in added),
                flags=page_flags,
            )
        )

    flags += builder.document_flags()
    if not builder.blocks:
        flags.append("no_content")
    return ParsedDocument(
        kind="pptx",
        parser="pptx/python-pptx",
        blocks=builder.blocks,
        pages=pages,
        flags=flags,
    )
